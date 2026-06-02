import sys, re, os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

def font(run, n='宋体', s=12, b=False, c=None):
    run.font.name = n; run.font.size = Pt(s); run.bold = b
    run.element.rPr.rFonts.set(qn('w:eastAsia'), n)
    if c: run.font.color.rgb = c

def ap(doc, text, sty='body'):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.space_before = Pt(2)
    if sty == 'title':
        p.paragraph_format.space_before = Pt(60); p.paragraph_format.space_after = Pt(20)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text); font(r, '黑体', 22, True); return p
    if sty == 'break':
        p.paragraph_format.space_before = Pt(16); p.paragraph_format.space_after = Pt(16)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run('· · ·'); font(r, '宋体', 10, c=RGBColor(180,180,180)); return p
    if sty == 'wechat':
        p.paragraph_format.left_indent = Cm(1)
        p.paragraph_format.space_after = Pt(1); p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.line_spacing = 1.2
        r = p.add_run(text); font(r, '宋体', 9.5, c=RGBColor(100,100,100)); return p
    if sty == 'code':
        p.paragraph_format.left_indent = Cm(0.8)
        p.paragraph_format.space_after = Pt(1); p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(text); font(r, 'Courier New', 8.5, c=RGBColor(80,80,80)); return p
    for part in re.split(r'(\*\*.*?\*\*)', text):
        if part.startswith('**') and part.endswith('**'):
            r = p.add_run(part[2:-2]); font(r, '宋体', 12, True)
        else:
            r = p.add_run(part); font(r, '宋体', 12)
    return p

ch = sys.argv[1]
fp = f'/Users/yupingwei/agent writing/chapters/chapter-{ch}.md'
with open(fp, 'r', encoding='utf-8') as f:
    content = f.read()

doc = Document()
for sec in doc.sections:
    sec.top_margin = Cm(2); sec.bottom_margin = Cm(2)
    sec.left_margin = Cm(2.2); sec.right_margin = Cm(2.2)
s = doc.styles['Normal']
s.font.name = '宋体'; s.font.size = Pt(12)
s.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
s.paragraph_format.line_spacing = 1.5

lines = content.split('\n')
i, in_code, buf = 0, False, []
while i < len(lines):
    l = lines[i]
    if l.strip().startswith('```'):
        if in_code:
            for cl in buf: ap(doc, cl, 'code')
            in_code = False; buf = []
        else: in_code = True
        i += 1; continue
    if in_code: buf.append(l); i += 1; continue
    if l.strip().startswith('# '): ap(doc, l.strip()[2:], 'title'); i += 1; continue
    if l.strip() == '---': ap(doc, '', 'break'); i += 1; continue
    if not l.strip(): i += 1; continue
    if re.match(r'^\d{2}:\d{2}\s+\S+[：:]', l.strip()): ap(doc, l.strip(), 'wechat'); i += 1; continue
    ap(doc, l); i += 1

names = {'01':'第一章-凌晨三点的工位','02':'第二章-万物有声','03':'第三章-杯子里的秘密'}
out = f"/Users/yupingwei/agent writing/word/{names[ch]}.docx"
doc.save(out)
print(f'✅ {os.path.basename(out)} ({round(os.path.getsize(out)/1024, 1)} KB)')
